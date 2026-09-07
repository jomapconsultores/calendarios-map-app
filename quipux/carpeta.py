# ------------------------------------------------------------
# Desarrollado por Marco Antonio Posligua San Martín
# ------------------------------------------------------------
"""Los oficios que se bajaron a mano, leídos como si los hubiera traído el robot.

El servidor no alcanza CuencaDOC y el acceso automático pide el texto de una
imagen cada vez que el Municipio decide pedirlo. Mientras eso siga así, la vía
corta es descargar los oficios desde el navegador, dejarlos en una carpeta y que
el programa los lea de ahí.

Lo que este módulo NO hace es una segunda versión del sistema. Saca de cada PDF
los mismos campos que el recolector saca de la ficha —número, asunto, quién lo
manda, fecha y plazo— y arma el mismo registro, para que a partir de ahí todo
siga por el camino de siempre: `deducir_plazo` calcula para cuándo es,
`almacen` lo guarda y `planificacion` lo lleva al cronograma. Dos puertas, una
sola casa.

Lo que se pierde respecto a la recolección automática, y conviene tener presente
al mirar el resultado:

  * el enlace al documento dentro de Quipux, que sólo existe si está impreso en
    el propio oficio;
  * la bandeja de la que venía, que en el sistema se sabe por dónde apareció y
    aquí no consta en ninguna parte;
  * la fecha de vencimiento QUE DA EL SISTEMA. Aquí el plazo sale de leer el
    texto, así que se marca como deducido —`seguro: False`— y en el cronograma
    entra con prioridad media en vez de alta. No es lo mismo que el sistema diga
    «vence el 12» a que una frase diga «en el término de cinco días», y quien lo
    mire tiene derecho a saber cuál de las dos está viendo.

Uso:
    python -m quipux.carpeta                     # lee ~/Dropbox/QUIPUX
    python -m quipux.carpeta --ver               # enseña lo que entendió, sin guardar
    python -m quipux.carpeta C:\\otra\\carpeta   # otra carpeta
"""
import os
import re
import sys
import zipfile
from datetime import date, datetime

from . import documentos as docs

CARPETA_POR_DEFECTO = os.path.join(os.path.expanduser('~'), 'Dropbox', 'QUIPUX')

MESES = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6,
         'julio': 7, 'agosto': 8, 'septiembre': 9, 'setiembre': 9, 'octubre': 10,
         'noviembre': 11, 'diciembre': 12}

# El número de un documento de Quipux: siglas de la entidad, año, secuencia y una
# letra al final que dice de qué tipo es (O de oficio, M de memorando…).
RE_NUMERO = re.compile(r'\b([A-Z][A-Z0-9]{2,}(?:-[A-Z0-9]+){1,6}-\d{4}-\d{3,6}-[A-Z])\b')
RE_FECHA_LARGA = re.compile(
    r'\b(\d{1,2})\s+de\s+([a-záéíóú]+)\s+(?:de[l]?\s+)?(\d{4})', re.I)
RE_ASUNTO = re.compile(r'asunto\s*:\s*(.+)', re.I)
RE_DE = re.compile(r'^\s*(?:de|remitente)\s*:\s*(.+)$', re.I | re.M)


def _texto_del_pdf(ruta):
    """Todo el texto del PDF. Si no se puede leer, cadena vacía y a seguir."""
    try:
        from pypdf import PdfReader
        lector = PdfReader(ruta)
        partes = []
        for pagina in lector.pages:
            try:
                partes.append(pagina.extract_text() or '')
            except Exception:
                continue
        return re.sub(r'[ \t]+', ' ', '\n'.join(partes)).strip()
    except Exception:
        return ''


def _texto_del_docx(ruta):
    """Los párrafos y también las tablas: en los oficios en Word, la mitad de lo
    que importa —fechas, responsables, montos— vive dentro de una tabla."""
    try:
        from docx import Document
        doc = Document(ruta)
        partes = [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila in tabla.rows:
                celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
                if celdas:
                    partes.append(' | '.join(celdas))
        return re.sub(r'[ \t]+', ' ', '\n'.join(partes)).strip()
    except Exception:
        return ''


def _filas_del_excel(ruta):
    """Las filas de la primera hoja, como listas de valores."""
    if ruta.lower().endswith('.xls'):
        try:
            import xlrd
            libro = xlrd.open_workbook(ruta)
            hoja = libro.sheet_by_index(0)
            return [[hoja.cell_value(f, c) for c in range(hoja.ncols)]
                    for f in range(hoja.nrows)]
        except Exception:
            return []
    try:
        from openpyxl import load_workbook
        libro = load_workbook(ruta, data_only=True, read_only=True)
        return [list(fila) for fila in libro[libro.sheetnames[0]].iter_rows(values_only=True)]
    except Exception:
        return []


def _texto_de_filas(filas, tope=400):
    return '\n'.join(' | '.join(str(v) for v in fila if v not in (None, ''))
                     for fila in filas[:tope])


def _fecha_de_celda(valor):
    """Una fecha de Excel puede venir como fecha de verdad o como texto."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return docs._fecha(valor) or _fecha_larga(str(valor or ''))


def _fecha_larga(texto):
    """«Cuenca, 6 de septiembre de 2026» -> date(2026, 9, 6)."""
    m = RE_FECHA_LARGA.search(texto or '')
    if not m:
        return None
    mes = MESES.get(m.group(2).lower().strip())
    if not mes:
        return None
    try:
        return date(int(m.group(3)), mes, int(m.group(1)))
    except ValueError:
        return None


def _primera_linea_util(texto, saltar):
    for linea in (texto or '').splitlines():
        limpia = linea.strip(' .:-')
        if len(limpia) > 12 and not any(p in limpia.lower() for p in saltar):
            return limpia
    return ''


# ============================================================
#  LAS MATRICES DE EXCEL
# ============================================================
# Un oficio es un documento; una matriz es una lista de compromisos. Tratar el
# Excel como un texto suelto dejaría veinte actividades convertidas en una sola
# línea, así que aquí cada fila con contenido se lee como lo que es: algo que
# alguien tiene que entregar, con su fecha.
COLUMNAS = {
    'asunto': ('actividad', 'tema', 'asunto', 'descripcion', 'descripción',
               'detalle', 'compromiso', 'tarea', 'producto', 'entregable',
               'accion', 'acción'),
    'plazo': ('plazo', 'fecha', 'vence', 'vencimiento', 'entrega', 'limite',
              'límite', 'cumplimiento', 'fin'),
    'responsable': ('responsable', 'encargado', 'asignado', 'a cargo'),
    'numero': ('numero', 'número', 'oficio', 'documento', 'codigo', 'código',
               'nro', 'n°'),
    'estado': ('estado', 'avance', 'situacion', 'situación'),
}


def _mapa_de_encabezados(filas, mirar=15):
    """Busca la fila de títulos y dice qué columna es qué.

    No se da por hecho que los títulos estén en la primera fila: las matrices
    del Municipio suelen abrir con el logo, el nombre de la dirección y un par
    de filas en blanco antes de empezar."""
    mejor, mejor_puntos, mapa_mejor = None, 0, {}
    for i, fila in enumerate(filas[:mirar]):
        mapa, puntos = {}, 0
        for col, valor in enumerate(fila):
            titulo = docs._norm(str(valor or ''))
            if not titulo:
                continue
            for campo, palabras in COLUMNAS.items():
                if campo in mapa:
                    continue
                if any(p in titulo for p in palabras):
                    mapa[campo] = col
                    puntos += 1
                    break
        # Con una sola columna reconocida no hay matriz que valga: sería
        # confundir un título cualquiera con una cabecera.
        if puntos >= 2 and puntos > mejor_puntos:
            mejor, mejor_puntos, mapa_mejor = i, puntos, mapa
    return mejor, mapa_mejor


def leer_matriz(ruta):
    """Un Excel -> una lista de registros, uno por fila con contenido.

    Si no se reconoce ninguna cabecera, devuelve lista vacía y quien llama lo
    trata como un documento más: un Excel puede ser una matriz de compromisos o
    puede ser cualquier otra cosa, y adivinar mal llenaría el cronograma de
    basura."""
    filas = _filas_del_excel(ruta)
    if not filas:
        return []
    cabecera, mapa = _mapa_de_encabezados(filas)
    if cabecera is None or 'asunto' not in mapa:
        return []

    nombre = os.path.basename(ruta)
    registros = []
    for i, fila in enumerate(filas[cabecera + 1:], start=cabecera + 2):
        def celda(campo):
            col = mapa.get(campo)
            if col is None or col >= len(fila):
                return ''
            return str(fila[col]).strip() if fila[col] not in (None, '') else ''

        asunto = celda('asunto')
        if len(asunto) < 4:
            continue

        plazo_col = mapa.get('plazo')
        fecha = None
        if plazo_col is not None and plazo_col < len(fila):
            fecha = _fecha_de_celda(fila[plazo_col])

        # El número lleva delante el nombre de la matriz. Sin eso, dos matrices
        # distintas con su fila «001» serían el mismo documento, y la segunda se
        # descartaría en silencio por repetida.
        base_archivo = os.path.splitext(nombre)[0][:34]
        propio = celda('numero')
        numero = f'{base_archivo}#{propio}' if propio else f'{base_archivo}#f{i}'
        registros.append({
            'id': numero,
            'numero': numero,
            'asunto': asunto[:300],
            'de': celda('responsable'),
            'fecha_doc': '',
            # De una matriz la fecha viene puesta por alguien, no deducida de
            # una frase: eso la hace tan buena como la del sistema.
            'vence': fecha.isoformat() if fecha else '',
            'area': '',
            'bandeja': f'Matriz · {nombre}',
            'enlace': '',
            'archivo': ruta,
            'texto': ' | '.join(str(v) for v in fila if v not in (None, '')),
            'estado_hoja': celda('estado'),
        })
    return registros


def leer_oficio(ruta):
    """Un PDF o un Word -> el mismo registro que devuelve la ficha del sistema."""
    texto = (_texto_del_docx(ruta) if ruta.lower().endswith(('.docx', '.docm'))
             else _texto_del_pdf(ruta))
    nombre = os.path.basename(ruta)

    # El número puede estar en el texto o en el propio nombre del archivo; el
    # del texto manda, porque el nombre lo pone quien descarga.
    m = RE_NUMERO.search(texto) or RE_NUMERO.search(nombre)
    numero = m.group(1) if m else os.path.splitext(nombre)[0][:60]

    asunto = ''
    m = RE_ASUNTO.search(texto)
    if m:
        asunto = m.group(1).strip()
        # El asunto de un oficio cabe en una línea; lo que sigue ya es el
        # cuerpo, y arrastrarlo entero deja títulos de mil caracteres.
        asunto = re.split(r'\s{3,}|\n', asunto)[0].strip()[:300]
    if not asunto:
        asunto = _primera_linea_util(
            texto, ('quipux', 'cuenca', 'gad ', 'municipal', 'oficio nro', 'señor'))[:300]

    m = RE_DE.search(texto)
    remitente = m.group(1).strip()[:120] if m else ''

    fecha_doc = _fecha_larga(texto)

    return {
        'id': numero,
        'numero': numero,
        'asunto': asunto or nombre,
        'de': remitente,
        'fecha_doc': fecha_doc.isoformat() if fecha_doc else '',
        'vence': '',          # el sistema lo daría; de un PDF no consta
        'area': '',           # se rellena con la carpeta que lo contiene
        'bandeja': 'Cargados a mano',
        'enlace': '',
        'archivo': ruta,
        'texto': texto,
    }


def _descomprimir(carpeta):
    """Deja abiertos los ZIP que haya, cada uno en su propia subcarpeta.

    Devuelve esas subcarpetas, no los nombres de los zip: hacen falta después
    para saber el área. Un zip llamado `lote_septiembre.zip` con
    `PLANIFICACION/` dentro no es un área llamada «lote septiembre» —es un
    envoltorio—, y sin esto todos sus documentos acababan en un proyecto con el
    nombre del archivo comprimido."""
    destinos = []
    for nombre in sorted(os.listdir(carpeta)):
        if not nombre.lower().endswith('.zip'):
            continue
        destino = os.path.join(carpeta, os.path.splitext(nombre)[0])
        if os.path.isdir(destino):
            destinos.append(destino)
            continue
        try:
            with zipfile.ZipFile(os.path.join(carpeta, nombre)) as z:
                z.extractall(destino)
            destinos.append(destino)
        except Exception as e:
            print(f'  no se pudo abrir {nombre}: {str(e)[:100]}')
    return destinos


def leer_carpeta(carpeta=None, registro=print):
    """Todos los oficios de la carpeta, con su plazo ya deducido.

    El área sale del nombre de la subcarpeta que los contiene: así, separar en
    «OBSERVATORIO» y «PLANIFICACION» al descargar es todo lo que hace falta para
    que cada documento acabe en el proyecto que le toca."""
    carpeta = carpeta or CARPETA_POR_DEFECTO
    if not os.path.isdir(carpeta):
        raise RuntimeError(f'No existe la carpeta {carpeta}')

    de_zip = _descomprimir(carpeta)
    if de_zip:
        registro(f'  {len(de_zip)} zip abierto(s)')

    encontrados = []
    for raiz, _, archivos in os.walk(carpeta):
        for nombre in sorted(archivos):
            if nombre.startswith('~$'):        # los temporales que deja Office
                continue
            if nombre.lower().endswith(('.pdf', '.docx', '.docm',
                                        '.xlsx', '.xlsm', '.xls')):
                encontrados.append(os.path.join(raiz, nombre))

    documentos, vistos = [], set()
    for ruta in encontrados:
        # Si el archivo salió de un zip, el área se cuenta desde dentro del zip:
        # su nombre es un envoltorio, no un área.
        raiz = carpeta
        for destino in de_zip:
            if ruta.startswith(destino + os.sep):
                raiz = destino
                break
        for doc in _leer_archivo(ruta):
            _completar(doc, ruta, raiz)
            # El mismo documento cargado dos veces es uno, no dos.
            if doc['id'] in vistos:
                registro(f"  repetido, se omite: {doc['id']}")
                continue
            vistos.add(doc['id'])
            documentos.append(doc)
    return documentos


def _leer_archivo(ruta):
    """Lo que haya dentro: un documento, o la lista de filas de una matriz."""
    if ruta.lower().endswith(('.xlsx', '.xlsm', '.xls')):
        # Un Excel puede traer veinte compromisos en veinte filas. Se intenta
        # leer como matriz; si no se le reconoce ninguna cabecera, se trata como
        # un documento más y su contenido se lee entero, que es mejor que
        # inventar veinte actividades a partir de una hoja que no lo es.
        filas = leer_matriz(ruta)
        if filas:
            return filas
        nombre = os.path.splitext(os.path.basename(ruta))[0]
        return [{
            'id': nombre[:60], 'numero': nombre[:60], 'asunto': nombre[:300],
            'de': '', 'fecha_doc': '', 'vence': '', 'area': '',
            'bandeja': 'Cargados a mano', 'enlace': '', 'archivo': ruta,
            'texto': _texto_de_filas(_filas_del_excel(ruta)),
        }]
    return [leer_oficio(ruta)]


def _completar(doc, ruta, carpeta):
    """Le pone el área y calcula para cuándo es."""
    # El área: la primera carpeta por debajo de la raíz, si la hay.
    relativa = os.path.relpath(os.path.dirname(ruta), carpeta)
    if relativa not in ('.', ''):
        doc['area'] = relativa.split(os.sep)[0].replace('_', ' ').strip()
    doc['area'] = doc['area'] or 'CuencaDOC'

    fecha, origen, seguro = docs.deducir_plazo(doc, doc.get('texto', ''))
    doc['plazo'] = {'fecha': fecha.isoformat() if fecha else '',
                    'origen': origen, 'seguro': seguro}
    doc['estado'] = 'abierto'
    doc['nuevo'] = True
    doc['n_adjuntos'] = 0
    return doc


def _resumen(documentos):
    con_plazo = [d for d in documentos if (d.get('plazo') or {}).get('fecha')]
    print('\n' + '=' * 60)
    print(f'Oficios leídos   : {len(documentos)}')
    print(f'Con plazo        : {len(con_plazo)}')
    sin_asunto = sum(1 for d in documentos if not d.get('asunto'))
    if sin_asunto:
        print(f'Sin asunto claro : {sin_asunto}  (habrá que mirarlos)')
    print('=' * 60)
    for d in documentos[:40]:
        plazo = (d.get('plazo') or {}).get('fecha') or '—'
        marca = '' if (d.get('plazo') or {}).get('seguro') else ' (deducido)'
        print(f"  {d['numero'][:34]:<34} {plazo:<11}{marca}")
        print(f"      {(d.get('asunto') or '')[:96]}")
        print(f"      área: {d.get('area', '')}"
              + (f" · responsable en la hoja: {d['de']}" if d.get('de') else ''))
    if len(documentos) > 40:
        print(f'  … y {len(documentos) - 40} más')


def main(argv=None):
    argv = argv if argv is not None else sys.argv[1:]
    solo_ver = '--ver' in argv
    resto = [a for a in argv if not a.startswith('--')]
    carpeta = resto[0] if resto else None

    documentos = leer_carpeta(carpeta)
    _resumen(documentos)

    if solo_ver:
        print('\n(--ver: no se guardó nada)')
        return 0

    from . import almacen, planificacion
    almacen.guardar(documentos)
    creadas, actualizadas = almacen.crear_tareas(documentos)
    print(f'\nEn el archivo local: {len(documentos)} documento(s)')
    print(f'Tareas             : {creadas} nueva(s), {actualizadas} al día')

    db = planificacion.cliente_de_la_plataforma()
    if db is None:
        print('Plataforma         : sin conexión; lo local queda guardado igual')
        return 0
    pub = planificacion.publicar(db, documentos)
    print(f"Publicados         : {pub.get('subidos', 0)}"
          + (f" (error: {pub['error']})" if pub.get('error') else ''))
    crono = planificacion.volcar(db, documentos)
    print(f"Cronograma         : {crono.get('creadas', 0)} nueva(s), "
          f"{crono.get('actualizadas', 0)} con plazo cambiado")
    return 0


if __name__ == '__main__':
    sys.exit(main())
